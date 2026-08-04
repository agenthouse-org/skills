#!/usr/bin/env python3
"""Create a professional role profile DOCX from the skill's JSON schema."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


TEXT = {
    "en": {
        "title": "Role Profile",
        "role": "Role",
        "purpose": "Purpose",
        "domains": "Domains",
        "requirements": "Requirements",
        "required": "Required",
        "preferred": "Preferred",
        "interfaces": "Interfaces",
        "exclusions": "Boundaries and exclusions",
        "assumptions": "Assumptions to validate",
        "matrix": "Responsibilities, authority, and obligations",
        "responsibility": "Responsibility",
        "info": "Info",
        "decision": "Decision",
        "direction": "Direction",
        "policy": "Policy",
        "control": "Control",
        "participation": "Participate",
        "authority_notes": "Authority notes",
        "quality": "Quality",
        "reporting": "Report",
        "approval": "Approval",
        "obligation_notes": "Obligation notes",
        "holder": "Role holder",
        "approved_by": "Approved by",
        "approval_date": "Approval date",
        "review_date": "Review date",
        "legend": "Authority and obligation definitions",
        "yes": "X",
    },
    "de": {
        "title": "Rollenbeschreibung",
        "role": "Rolle",
        "purpose": "Zweck",
        "domains": "Domänen",
        "requirements": "Anforderungsprofil",
        "required": "Erforderlich",
        "preferred": "Wünschenswert",
        "interfaces": "Schnittstellen",
        "exclusions": "Abgrenzungen und Einschränkungen",
        "assumptions": "Zu validierende Annahmen",
        "matrix": "Verantwortung, Befugnisse und Pflichten",
        "responsibility": "Verantwortungsbereich",
        "info": "Information",
        "decision": "Entscheidung",
        "direction": "Fachl. Weisung",
        "policy": "Richtlinie",
        "control": "Kontrolle",
        "participation": "Mitwirkung",
        "authority_notes": "Befugnisnotizen",
        "quality": "Qualität",
        "reporting": "Bericht",
        "approval": "Freigabe",
        "obligation_notes": "Pflichtnotizen",
        "holder": "Rolleninhaber",
        "approved_by": "Genehmigt durch",
        "approval_date": "Genehmigungsdatum",
        "review_date": "Prüfdatum",
        "legend": "Definitionen der Befugnisse und Pflichten",
        "yes": "X",
    },
}

DEFINITIONS = {
    "en": [
        ("Information authority", "Request and receive information relevant to the responsibility."),
        ("Decision authority", "Make decisions and grant approvals within the stated boundary."),
        ("Functional-direction authority", "Direct named roles on professional or process matters."),
        ("Policy authority", "Create, maintain, or change rules and standards in scope."),
        ("Control authority", "Inspect and evaluate quality or compliance."),
        ("Participation authority", "Advise and contribute to another role's decision."),
        ("Quality-control obligation", "Ensure the defined quality or compliance level."),
        ("Reporting obligation", "Provide status or a defined report to a named recipient."),
        ("Approval obligation", "Obtain another role's approval before acting."),
    ],
    "de": [
        ("Informationsbefugnis", "Informationen zur Tätigkeit anfragen und erhalten."),
        ("Entscheidungsbefugnis", "Im definierten Rahmen entscheiden und genehmigen."),
        ("Fachliche Weisungsbefugnis", "Genannten Rollen fachliche oder prozessuale Anordnungen erteilen."),
        ("Richtlinienbefugnis", "Regeln und Standards im Geltungsbereich erstellen und ändern."),
        ("Kontrollbefugnis", "Qualität oder Regelkonformität prüfen und bewerten."),
        ("Mitwirkungsbefugnis", "Eine andere entscheidende Rolle beraten und unterstützen."),
        ("Kontrollpflicht", "Definierte Qualität oder Regelkonformität sicherstellen."),
        ("Berichtspflicht", "Status oder Bericht an einen definierten Empfänger liefern."),
        ("Freigabepflicht", "Vor dem Handeln die Freigabe einer anderen Rolle einholen."),
    ],
}

BLUE = "17365D"
TEAL = "1F6D72"
LIGHT = "EAF1F5"
PALE = "F4F7F9"
WHITE = "FFFFFF"


def shade(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def set_cell_margin(cell, top=80, start=90, bottom=80, end=90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, size=9, bold=False, color=None) -> None:
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def write_cell(cell, text, *, bold=False, color=None, size=8, align=None) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text or ""))
    set_font(run, size=size, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margin(cell)


def add_bullets(cell, items, prefix=None) -> None:
    cell.text = ""
    if prefix:
        p = cell.paragraphs[0]
        r = p.add_run(prefix)
        set_font(r, size=9, bold=True, color=BLUE)
    for item in items or []:
        p = cell.add_paragraph(style=None)
        p.style = "List Bullet"
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(str(item))
        set_font(r, size=9)


def add_section_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(9)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_font(r, size=13, bold=True, color=BLUE)


def build(input_path: Path, output_path: Path) -> None:
    data = json.loads(input_path.read_text(encoding="utf-8"))
    language = data.get("language", "en")
    language = language if language in TEXT else "en"
    t = TEXT[language]

    if not data.get("role_name") or not data.get("purpose"):
        raise ValueError("role_name and purpose are required")
    if not data.get("responsibilities"):
        raise ValueError("at least one responsibility is required")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)
    section.left_margin = Cm(1.7)
    section.right_margin = Cm(1.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(t["title"].upper())
    set_font(r, size=10, bold=True, color=TEAL)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(data["role_name"])
    set_font(r, size=22, bold=True, color=BLUE)

    overview = doc.add_table(rows=4, cols=2)
    overview.alignment = WD_TABLE_ALIGNMENT.CENTER
    overview.autofit = False
    overview.columns[0].width = Cm(3.6)
    overview.columns[1].width = Cm(13.6)
    fields = [
        (t["purpose"], data.get("purpose", "")),
        (t["domains"], "\n".join(f"• {x}" for x in data.get("domains", []))),
        (t["interfaces"], "\n".join(f"• {x}" for x in data.get("interfaces", []))),
        (t["requirements"], ""),
    ]
    for i, (label, value) in enumerate(fields):
        write_cell(overview.cell(i, 0), label, bold=True, color=WHITE, size=9)
        shade(overview.cell(i, 0), BLUE)
        if i < 3:
            write_cell(overview.cell(i, 1), value, size=9)
            shade(overview.cell(i, 1), PALE if i % 2 else WHITE)
        else:
            req = data.get("requirements", {})
            add_bullets(overview.cell(i, 1), req.get("required", []), f'{t["required"]}:')
            if req.get("preferred"):
                p2 = overview.cell(i, 1).add_paragraph()
                r2 = p2.add_run(f'{t["preferred"]}:')
                set_font(r2, size=9, bold=True, color=BLUE)
                for item in req.get("preferred", []):
                    p3 = overview.cell(i, 1).add_paragraph(style="List Bullet")
                    r3 = p3.add_run(str(item))
                    set_font(r3, size=9)

    matrix_section = doc.add_section()
    matrix_section.orientation = WD_ORIENT.LANDSCAPE
    matrix_section.page_width, matrix_section.page_height = section.page_height, section.page_width
    matrix_section.top_margin = Cm(1.3)
    matrix_section.bottom_margin = Cm(1.3)
    matrix_section.left_margin = Cm(1.2)
    matrix_section.right_margin = Cm(1.2)

    add_section_heading(doc, t["matrix"])

    headers = [
        t["responsibility"], t["info"], t["decision"], t["direction"], t["policy"],
        t["control"], t["participation"], t["authority_notes"], t["quality"],
        t["reporting"], t["approval"], t["obligation_notes"],
    ]
    rows = data["responsibilities"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [5.0, 1.05, 1.2, 1.25, 1.05, 1.05, 1.2, 3.5, 1.05, 1.05, 1.05, 3.5]
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        write_cell(cell, header, bold=True, color=WHITE, size=7, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(cell, BLUE if i in (0, 7, 11) else TEAL)
        cell.width = Cm(widths[i])

    authority_keys = ["information", "decision", "functional_direction", "policy", "control", "participation"]
    obligation_keys = ["quality_control", "reporting", "approval"]
    for row_index, item in enumerate(rows):
        cells = table.add_row().cells
        values = [item.get("responsibility", "")]
        values.extend(t["yes"] if item.get("authorities", {}).get(k) else "" for k in authority_keys)
        values.append(item.get("authority_notes", ""))
        values.extend(t["yes"] if item.get("obligations", {}).get(k) else "" for k in obligation_keys)
        values.append(item.get("obligation_notes", ""))
        for i, value in enumerate(values):
            align = WD_ALIGN_PARAGRAPH.CENTER if i in (1, 2, 3, 4, 5, 6, 8, 9, 10) else WD_ALIGN_PARAGRAPH.LEFT
            write_cell(cells[i], value, size=7, align=align)
            cells[i].width = Cm(widths[i])
            if row_index % 2:
                shade(cells[i], PALE)

    add_section_heading(doc, t["exclusions"])
    for item in data.get("exclusions", []):
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(str(item))
        set_font(r, size=9)

    if data.get("assumptions"):
        add_section_heading(doc, t["assumptions"])
        for item in data["assumptions"]:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(str(item))
            set_font(r, size=9, color="9C5700")

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    labels = [t["holder"], t["approved_by"], t["approval_date"], t["review_date"]]
    values = [data.get("holder", ""), data.get("approved_by", ""), data.get("approval_date", ""), data.get("review_date", "")]
    for i in range(4):
        write_cell(meta.cell(0, i), labels[i], bold=True, color=WHITE, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(meta.cell(0, i), BLUE)
        write_cell(meta.cell(1, i), values[i], size=8, align=WD_ALIGN_PARAGRAPH.CENTER)

    add_section_heading(doc, t["legend"])
    legend = doc.add_table(rows=0, cols=2)
    legend.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (term, definition) in enumerate(DEFINITIONS[language]):
        cells = legend.add_row().cells
        write_cell(cells[0], term, bold=True, color=BLUE, size=8)
        write_cell(cells[1], definition, size=8)
        if i % 2:
            shade(cells[0], PALE)
            shade(cells[1], PALE)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: create_role_docx.py INPUT.json OUTPUT.docx")
    build(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
